"""Lectura de los crudos (Excel o CSV) y canonicalización de celdas al formato de salida.

Los Excel se leen con openpyxl (valores calculados, columnas ocultas incluidas) y los
CSV de portal con el módulo `csv` (respetando comillas). En ambos casos cada celda se
convierte a su representación final de texto ANTES de operar — fechas DD/MM/AAAA y
números es-AR — así el resto del motor trabaja siempre sobre strings
estables y el escritor no tiene que decidir nada.
"""

import csv
import datetime as dt
from pathlib import Path

import openpyxl
import polars as pl

EXTENSIONES_EXCEL = {".xlsx", ".xlsm"}
EXTENSIONES_CSV = {".csv", ".txt"}
EXTENSIONES_SOPORTADAS = EXTENSIONES_EXCEL | EXTENSIONES_CSV


def leer_crudo(ruta: Path, entrada_cfg: dict, formato_cfg: dict):
    """Devuelve (nombres_salida, df) — nombres tal cual el encabezado original
    (pueden repetirse) y un DataFrame polars todo Utf8 con nombres internos únicos."""
    ruta = Path(ruta)
    ext = ruta.suffix.lower()
    if ext in EXTENSIONES_EXCEL:
        filas = _filas_excel(ruta, entrada_cfg)
    elif ext in EXTENSIONES_CSV:
        filas = _filas_csv(ruta, entrada_cfg)
    else:
        raise ValueError(
            f"Extensión no soportada: '{ext}'. El motor lee {sorted(EXTENSIONES_SOPORTADAS)}."
        )

    fila_encabezado = _fila_encabezado(filas, entrada_cfg)
    if len(filas) < fila_encabezado:
        raise ValueError(f"El archivo tiene {len(filas)} filas; el encabezado se esperaba en la fila {fila_encabezado}")
    omitir = _columnas_a_omitir(filas[fila_encabezado - 1], entrada_cfg)

    encabezado = list(filas[fila_encabezado - 1])[omitir:]
    # ancho útil = hasta el último título no vacío (recorta colas de columnas fantasma)
    ancho = max((i + 1 for i, v in enumerate(encabezado) if v is not None and str(v).strip() != ""), default=0)
    if ancho == 0:
        raise ValueError(f"La fila {fila_encabezado} no tiene títulos — revisar 'fila_encabezado' en la config")
    nombres = [str(v) if v is not None else "" for v in encabezado[:ancho]]
    vacios = [i for i, n in enumerate(nombres) if n.strip() == ""]
    if vacios and not entrada_cfg.get("permitir_titulos_vacios", False):
        raise ValueError(
            f"Títulos vacíos en las posiciones {vacios} del encabezado — "
            "si son columnas descartables al inicio, usar 'columnas_omitir_iniciales'; "
            "si el layout del agente realmente trae títulos en blanco, "
            "declarar 'permitir_titulos_vacios: true' en la config"
        )

    datos = [list(f)[omitir:omitir + ancho] + [None] * max(0, ancho - (len(f) - omitir))
             for f in filas[fila_encabezado:]]
    # descartar filas totalmente vacías (colas típicas de Excel)
    datos = [f for f in datos if any(v is not None and str(v).strip() != "" for v in f)]

    fechas = _indices_fecha(nombres, formato_cfg)
    columnas = {}
    for j, nombre in enumerate(nombres):
        interno = f"c{j}"  # nombres internos únicos; los reales salen en el escritor
        es_fecha = j in fechas
        columnas[interno] = [canonicalizar(f[j], nombre, formato_cfg, es_fecha) for f in datos]
    df = pl.DataFrame(columnas)
    return nombres, df


def _filas_excel(ruta: Path, entrada_cfg: dict):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = entrada_cfg.get("hoja", 0)
        ws = wb[hoja] if isinstance(hoja, str) else wb[wb.sheetnames[hoja]]
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _filas_csv(ruta: Path, entrada_cfg: dict):
    """CSV de portal: valores ya son texto, pero pueden venir entrecomillados y con
    saltos de línea adentro — el módulo `csv` los resuelve; canonicalizar() después
    aplana los saltos para que no rompan el CSV sin comillas de la salida."""
    encoding = entrada_cfg.get("encoding", "utf-8-sig")  # los portales suelen mandar BOM
    sep = entrada_cfg.get("separador_origen", ",")
    with open(ruta, encoding=encoding, newline="") as f:
        return [tuple(fila) for fila in csv.reader(f, delimiter=sep)]


def _fila_encabezado(filas, entrada_cfg) -> int:
    """Resuelve `fila_encabezado`, que puede venir de dos formas.

    - Un número (1-based): el caso normal, el encabezado está siempre igual.
    - `{ancla: "HAWB", buscar_hasta: 40}`: se busca la primera fila que contenga
      ese título exacto. Hay reportes que anteponen un preámbulo con los filtros
      aplicados (empresa, modo de transporte, consignatario, …) y su largo depende
      de cuántos filtros se pusieron al bajarlo: el encabezado puede estar en la
      fila 1 un día y en la 23 al siguiente. Clavar el número rompe el agente cada
      vez que alguien cambia la consulta.
    """
    cfg = entrada_cfg.get("fila_encabezado", 1)
    if not isinstance(cfg, dict):
        return int(cfg)

    ancla = str(cfg["ancla"]).strip()
    limite = min(len(filas), int(cfg.get("buscar_hasta", 40)))
    for i in range(limite):
        if any(v is not None and str(v).strip() == ancla for v in filas[i]):
            return i + 1
    raise ValueError(
        f"No se encontró la fila de encabezado: ningún título '{ancla}' en las "
        f"primeras {limite} filas. ¿Cambió el layout del agente o el ancla?"
    )


def _columnas_a_omitir(encabezado, entrada_cfg) -> int:
    """Resuelve `columnas_omitir_iniciales`, que puede venir de dos formas.

    - Un número: se descartan siempre esas columnas del principio.
    - `{hasta: "HAWB"}`: se descarta todo lo que está a la izquierda de ese título,
      sean 0, 1 o 2 columnas, vacías o con datos. Hay reportes que a veces traen dos
      columnas vacías adelante (A y B) y a veces no: con un 2 fijo, el día que no
      vienen se pierden las dos primeras columnas reales. Saltear "solo las sin
      título" tampoco alcanza: si A y B llegaran con título, el CSV saldría con dos
      columnas de más, y el destino lee por posición.
    """
    cfg = entrada_cfg.get("columnas_omitir_iniciales", 0)
    if not isinstance(cfg, dict):
        return int(cfg)
    hasta = str(cfg["hasta"]).strip()
    titulos = [str(v).strip() if v is not None else "" for v in encabezado]
    if hasta not in titulos:
        raise ValueError(
            f"columnas_omitir_iniciales: no está el título '{hasta}' en la fila de "
            f"encabezado. ¿Cambió el layout del agente?")
    omitir = titulos.index(hasta)
    print(f"  Lector: {omitir} columnas antes de '{hasta}' se descartan")
    return omitir


def _es_columna_fecha(nombre_col: str, formato_cfg: dict) -> bool:
    """Comparación tolerante a espacios: los portales exportan títulos con espacios
    de más ('Master Instructions to SSL ') y el resto del motor ya los tolera."""
    return nombre_col.strip() in {
        str(c["columna"] if isinstance(c, dict) else c).strip()
        for c in formato_cfg.get("columnas_fecha", [])
    }


def _indices_fecha(nombres, formato_cfg):
    """Resuelve `columnas_fecha` a posiciones concretas del encabezado.

    Normalmente alcanza con el nombre, pero un crudo puede repetir un título y
    querer la conversión en una sola de las dos (un reporte con dos
    `Loaded for pre-carriage`, donde el destino espera la primera como fecha y la
    segunda con el texto del portal, '06.03.2026 00:00'). Para ese caso la config escribe `{columna: X, ocurrencia: N}` (1-based).
    """
    limpios = [n.strip() for n in nombres]
    indices = set()
    for entrada in formato_cfg.get("columnas_fecha", []) or []:
        if isinstance(entrada, dict):
            nombre, ocurrencia = str(entrada["columna"]).strip(), entrada.get("ocurrencia")
        else:
            nombre, ocurrencia = str(entrada).strip(), None
        posiciones = [i for i, n in enumerate(limpios) if n == nombre]
        if ocurrencia is None:
            indices.update(posiciones)
        elif 1 <= ocurrencia <= len(posiciones):
            indices.add(posiciones[ocurrencia - 1])
        else:
            raise ValueError(
                f"columnas_fecha: se pidió la ocurrencia {ocurrencia} de '{nombre}', "
                f"pero aparece {len(posiciones)} vez/veces en el encabezado"
            )
    return indices


def canonicalizar(valor, nombre_col: str, formato_cfg: dict, es_fecha=None) -> str:
    """`es_fecha` pisa la detección por nombre (ver `_indices_fecha`); si va en
    None se resuelve por el título, que es el caso normal."""
    if valor is None:
        return ""
    columna_fecha = (es_fecha if es_fecha is not None
                     else _es_columna_fecha(nombre_col, formato_cfg))
    if isinstance(valor, dt.datetime):
        return valor.strftime(formato_cfg["fecha"])
    if isinstance(valor, dt.date):
        return valor.strftime(formato_cfg["fecha"])
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, (int, float)):
        # en columnas declaradas de fecha, un número es un serial de Excel crudo
        # (celda sin formato en el origen): 42097 -> 03/04/2015
        if columna_fecha:
            fecha = dt.date(1899, 12, 30) + dt.timedelta(days=int(valor))
            return fecha.strftime(formato_cfg["fecha"])
        spec = formato_cfg["numeros"].get(nombre_col, {})
        return formatear_numero(valor, spec.get("decimales"), spec.get("miles", False))
    # saltos de línea dentro de la celda (alt-Enter de Excel) romperían el CSV
    # sin comillas de la salida: se aplanan a espacio, sin tocar el resto del texto
    texto = str(valor).replace("\r\n", " ").replace("\r", " ").replace("\n", " ")

    # Fechas que llegan como TEXTO (típico de los CSV de portal: "2026-04-02").
    # No alcanza con cambiar '-' por '/': eso daría 2026/04/02 en vez de
    # 02/04/2026. Hay que parsear y reformatear.
    # `fecha_origen` acepta una lista: un mismo reporte puede mezclar formatos
    # ("27-Apr-26 10:48 AM" en una columna y "04-May-26" en otra).
    origen = formato_cfg.get("fecha_origen")
    if origen and texto.strip() and columna_fecha:
        formatos = [origen] if isinstance(origen, str) else list(origen)
        for fmt in formatos:
            try:
                return dt.datetime.strptime(texto.strip(), fmt).strftime(formato_cfg["fecha"])
            except ValueError:
                continue  # probar el siguiente formato declarado
        # ninguno matcheó: no es una fecha (texto libre, celda rara) y se deja crudo
    return texto


def formatear_numero(v, decimales=None, miles=False) -> str:
    """Números al estilo es-AR: coma decimal y, si se pide, punto de miles."""
    if decimales is None:
        if float(v).is_integer():
            s = f"{int(v):,}" if miles else str(int(v))
        else:
            s = f"{v:,}" if miles else str(v)
    else:
        s = f"{float(v):,.{decimales}f}"
        if not miles:
            s = s.replace(",", "")
    # swap separadores: 1,234.56 -> 1.234,56
    return s.replace(",", "\x00").replace(".", ",").replace("\x00", ".")
