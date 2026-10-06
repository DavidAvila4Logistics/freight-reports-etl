"""Correcciones de fila sobre el CSV final (`data/correcciones/`) y su paso por el registro.

El agente de prueba es el caso difícil: su clave son DOS columnas (`HAWB` +
`Container`) y las dos se pueden corregir, así que es donde de verdad hace falta
`_match`. Cierra con una columna sin título, como algunos CSV de destino.
"""

import datetime as dt
import os

import openpyxl
import pytest

from conftest import df_de
from motor import correcciones
from motor.registro import upsert
import main as motor_main

NOMBRES = ["HAWB", "Customer References", "Container", "Vessel", "ETA", ""]
CLAVE = ["HAWB", "Container"]
SEP, ENC = ";", "cp1252"

# El reporte de hoy, ya pasado por el pipeline: la ventana de días YA filtró.
PORTAL = [
    ["S001", "701000001", "ABCU1000001", "SEA STAR", "10/09/2026", ""],
    ["S002", "701000002", "ABCU1000002", "OCEAN PEARL", "11/09/2026", ""],
    ["S003", "", "ABCU1000003", "OCEAN PEARL", "12/09/2026", ""],
]


def armar(filas=None):
    return df_de(NOMBRES, PORTAL if filas is None else filas)[1]


def escribir(carpeta, filas, cabecera=None, nombre="01-TEST.csv"):
    """Un archivo de correcciones con el contrato real: cp1252, ';', CRLF."""
    ruta = carpeta / nombre
    cab = NOMBRES if cabecera is None else cabecera
    texto = SEP.join(cab) + "\r\n" + "".join(SEP.join(f) + "\r\n" for f in filas)
    ruta.write_bytes(texto.encode(ENC, errors="replace"))
    return ruta


def escribir_xlsx(carpeta, filas, nombre="01-TEST.xlsx"):
    """Como guarda Excel un .xlsx: con tipos de celda (fecha, número), no strings."""
    ruta = carpeta / nombre
    wb = openpyxl.Workbook()
    wb.active.append(NOMBRES)
    for f in filas:
        wb.active.append(f)
    wb.save(ruta)
    return ruta


def con_mtime(ruta, cuando):
    os.utime(ruta, (cuando, cuando))
    return ruta


def correr(ruta, df=None, clave=None):
    """(df, tocadas, renombradas, avisos) — el camino completo cargar + aplicar."""
    df = armar() if df is None else df
    avisos = []
    filas = correcciones.cargar(ruta, ENC, SEP, NOMBRES, avisos)
    df, tocadas, renombradas = correcciones.aplicar(NOMBRES, df, filas, clave or CLAVE, avisos)
    return df, tocadas, renombradas, avisos


def tipos(avisos):
    return sorted(a["tipo"] for a in avisos)


def filas_de(df):
    return [list(f) for f in df.iter_rows()]


def test_pisa_la_celda_declarada_y_no_toca_el_resto(tmp_path):
    ruta = escribir(tmp_path, [["S001", "", "ABCU1000001", "SEA BREEZE", "", ""]])
    df, tocadas, renombradas, avisos = correr(ruta)
    assert filas_de(df) == [["S001", "701000001", "ABCU1000001", "SEA BREEZE", "10/09/2026", ""],
                            PORTAL[1], PORTAL[2]]
    assert tocadas == {"S001|ABCU1000001"} and renombradas == {}
    assert tipos(avisos) == ["correccion_fila"]


def test_celda_vacia_no_toca_y_el_centinela_blanquea(tmp_path):
    ruta = escribir(tmp_path, [["S002", "<VACIO>", "ABCU1000002", "", "", ""]])
    df, _, _, _ = correr(ruta)
    assert filas_de(df)[1] == ["S002", "", "ABCU1000002", "OCEAN PEARL", "11/09/2026", ""]


def test_nunca_agrega_filas(tmp_path):
    ruta = escribir(tmp_path, [["S999", "", "ABCU9999999", "OTRO BUQUE", "", ""]])
    df, tocadas, _, avisos = correr(ruta)
    assert filas_de(df) == PORTAL and not tocadas
    assert tipos(avisos) == ["correccion_sin_destino"]


def test_no_revive_una_carga_que_la_ventana_dejo_afuera(tmp_path):
    ruta = escribir(tmp_path, [["S001", "", "ABCU1000001", "SEA BREEZE", "", ""]])
    df, tocadas, _, avisos = correr(ruta, df=armar(PORTAL[1:]))
    assert df.height == 2 and not tocadas
    assert "correccion_sin_destino" in tipos(avisos)


def test_dos_correcciones_para_la_misma_carga_no_se_aplica_ninguna(tmp_path):
    ruta = escribir(tmp_path, [["S001", "", "ABCU1000001", "BUQUE A", "", ""],
                               ["S001", "", "ABCU1000001", "BUQUE B", "", ""]])
    df, _, _, avisos = correr(ruta)
    assert filas_de(df) == PORTAL and tipos(avisos) == ["correccion_ambigua"]


def test_encabezado_desalineado_rebota_como_error(tmp_path):
    movido = ["HAWB", "Container", "Customer References", "Vessel", "ETA", ""]
    ruta = escribir(tmp_path, [["S001", "ABCU1000001", "", "BUQUE A", "", ""]], cabecera=movido)
    df, _, _, avisos = correr(ruta)
    assert filas_de(df) == PORTAL and [a["nivel"] for a in avisos] == ["error"]
    # El mismo encabezado viejo, pero sin correcciones cargadas, no molesta a nadie.
    _, _, _, avisos = correr(escribir(tmp_path, [], cabecera=movido))
    assert avisos == []


@pytest.mark.parametrize("cabecera, fila", [
    (NOMBRES[:-1], ["S001", "", "ABCU1000001", "BUQUE A"]),                     # sin el ';' final
    (NOMBRES + [""], ["S001", "", "ABCU1000001", "BUQUE A", "", "", ""]),       # con uno de más
])
def test_tolera_la_terminadora_vacia(tmp_path, cabecera, fila):
    df, _, _, avisos = correr(escribir(tmp_path, [fila], cabecera=cabecera))
    assert filas_de(df)[0][3] == "BUQUE A" and tipos(avisos) == ["correccion_fila"]


def test_match_corrige_una_columna_de_la_clave(tmp_path):
    ruta = escribir(tmp_path, [["", "", "ABCU7654321", "", "", "", "HAWB=S001"]],
                    cabecera=NOMBRES + ["_match"])
    df, tocadas, renombradas, avisos = correr(ruta)
    assert filas_de(df)[0] == ["S001", "701000001", "ABCU7654321", "SEA STAR", "10/09/2026", ""]
    assert renombradas == {"S001|ABCU1000001": "S001|ABCU7654321"}
    assert tocadas == {"S001|ABCU7654321"}
    assert "correccion_identidad_destino" in tipos(avisos)


def test_match_con_varios_pares(tmp_path):
    ruta = escribir(tmp_path, [["S009", "", "", "", "", "", "HAWB=S003|Vessel=OCEAN PEARL"]],
                    cabecera=NOMBRES + ["_match"])
    df, _, renombradas, _ = correr(ruta)
    assert filas_de(df)[2][0] == "S009"
    assert renombradas == {"S003|ABCU1000003": "S009|ABCU1000003"}


def test_match_ambiguo_que_toca_la_clave_no_se_aplica(tmp_path):
    ruta = escribir(tmp_path, [["", "", "ABCU0000000", "", "", "", "Vessel=OCEAN PEARL"]],
                    cabecera=NOMBRES + ["_match"])
    df, _, _, avisos = correr(ruta)
    assert filas_de(df) == PORTAL and tipos(avisos) == ["correccion_ambigua"]


def test_match_a_varias_filas_sin_tocar_la_clave_se_aplica_a_todas(tmp_path):
    # Una factura con cientos de piezas: corregirles el ETA tiene que ser una fila.
    ruta = escribir(tmp_path, [["", "", "", "", "20/09/2026", "", "Vessel=OCEAN PEARL"]],
                    cabecera=NOMBRES + ["_match"])
    df, tocadas, renombradas, avisos = correr(ruta)
    assert [f[4] for f in filas_de(df)] == ["10/09/2026", "20/09/2026", "20/09/2026"]
    assert tocadas == {"S002|ABCU1000002", "S003|ABCU1000003"} and renombradas == {}
    assert tipos(avisos) == ["correccion_fila"] and "2 filas" in avisos[0]["mensaje"]


def test_clave_incompleta_sin_match_no_se_aplica(tmp_path):
    df, _, _, avisos = correr(escribir(tmp_path, [["S001", "", "", "BUQUE A", "", ""]]))
    assert filas_de(df) == PORTAL and tipos(avisos) == ["correccion_clave_incompleta"]


def test_columnas_de_control(tmp_path):
    ruta = escribir(tmp_path, [["S001", "", "ABCU1000001", "BUQUE A", "", "", "",
                                "lo confirmó la naviera por mail"]],
                    cabecera=NOMBRES + ["_match", "_motivo"])
    _, _, _, avisos = correr(ruta)
    assert any("naviera" in a["mensaje"] for a in avisos)
    ruta = escribir(tmp_path, [["S001", "", "ABCU1000001", "BUQUE A", "", "", "x"]],
                    cabecera=NOMBRES + ["_loquesea"])
    _, _, _, avisos = correr(ruta)
    assert "correccion_control_desconocida" in tipos(avisos)


def test_sin_archivo_no_pasa_nada(tmp_path):
    assert correr(tmp_path / "no_existe.csv")[3] == []


def test_main_aplica_las_correcciones_del_agente(tmp_path):
    escribir(tmp_path, [["S001", "", "ABCU1000001", "BUQUE DEL MAIN", "", ""]])
    cfg = {"salida": {"nombre": "01-TEST.csv", "separador": SEP, "encoding": ENC},
           "formato": {"fecha": "%d/%m/%Y"}, "registro": {"clave": CLAVE}}
    cfg_global = {"dir_correcciones": str(tmp_path), "dir_backup": str(tmp_path / "backup")}
    df, _, _ = motor_main.aplicar_correcciones(cfg, cfg_global, NOMBRES, armar(), {})
    assert filas_de(df)[0][3] == "BUQUE DEL MAIN"
    # Sin `registro.clave` no hay con qué identificar la carga: avisa y no aplica.
    aux = {}
    df, _, _ = motor_main.aplicar_correcciones(dict(cfg, registro={}), cfg_global, NOMBRES,
                                               armar(), aux)
    assert filas_de(df) == PORTAL and tipos(aux["_avisos"]) == ["correccion_sin_clave"]


# ── El registro rebautiza la carga, no la duplica ───────────────────────────
def test_registro_con_clave_corregida():
    cfg_reg = {"clave": CLAVE}
    cols, df_reg, _ = upsert(NOMBRES, armar(), cfg_reg, [], [], dt.date(2026, 9, 1), "portal.xlsx")
    previas = [dict(zip(cols, f)) for f in df_reg.iter_rows()]
    assert df_reg.height == 3

    corregido = [["S001", "701000001", "ABCU7654321", "SEA STAR", "10/09/2026", ""],
                 PORTAL[1], PORTAL[2]]
    manana = dt.date(2026, 9, 2)
    _, sin_mapa, _ = upsert(NOMBRES, armar(corregido), cfg_reg, cols, previas, manana, "p.xlsx")
    assert sin_mapa.height == 4, "sin el mapa de renombre la carga se duplica"

    cols2, con_mapa, resumen = upsert(
        NOMBRES, armar(corregido), cfg_reg, cols, previas, manana, "p.xlsx",
        corregidas={"S001|ABCU7654321"}, renombradas={"S001|ABCU1000001": "S001|ABCU7654321"})
    assert con_mapa.height == 3
    tocada = [dict(zip(cols2, f)) for f in con_mapa.iter_rows()][0]
    assert tocada["_clave"] == "S001|ABCU7654321"
    assert tocada["_primera_carga"] == "01/09/2026"           # conserva su historia
    assert (tocada["_estado"], tocada["_corregido_por"]) == ("corregido", "correcciones")
    assert resumen["corregidos"] == 1


# ── Un .csv de correcciones guardado como .xlsx desde Excel ─────────────────
T0 = dt.datetime(2026, 9, 7, 10, 0, 0).timestamp()


def test_xlsx_mas_nuevo_se_convierte_y_reemplaza_al_csv(tmp_path):
    backup = tmp_path / "backup"
    ruta_csv = con_mtime(escribir(tmp_path, []), T0)
    ruta_xlsx = con_mtime(escribir_xlsx(
        tmp_path, [["S001", None, "ABCU1000001", "OCEAN PEARL", dt.date(2026, 9, 15), None]]), T0 + 10)
    avisos = []
    assert correcciones.resolver_excel_pendiente(ruta_csv, backup, SEP, ENC, "%d/%m/%Y", avisos)
    assert not ruta_xlsx.exists()
    texto = ruta_csv.read_text(encoding=ENC)
    assert "OCEAN PEARL" in texto and "15/09/2026" in texto
    assert list(backup.glob("*.csv")) and list(backup.glob("*.xlsx"))
    assert "correccion_excel_convertido" in tipos(avisos)
    assert "S001|ABCU1000001" in correr(ruta_csv)[1]


def test_xlsx_mas_viejo_no_pisa_al_csv(tmp_path):
    ruta_csv = con_mtime(escribir(tmp_path, [["S002", "", "ABCU1000002", "NUEVO", "", ""]]), T0 + 100)
    ruta_xlsx = con_mtime(escribir_xlsx(tmp_path, [["S002", "", "ABCU1000002", "VIEJO", "", ""]]), T0)
    antes, avisos = ruta_csv.read_bytes(), []
    assert not correcciones.resolver_excel_pendiente(ruta_csv, tmp_path / "b", SEP, ENC, "%d/%m/%Y",
                                                     avisos)
    assert ruta_csv.read_bytes() == antes and ruta_xlsx.exists()
    assert tipos(avisos) == ["correccion_excel_viejo"]


def test_numero_tipeado_en_excel_sale_sin_punto_cero(tmp_path):
    ruta_csv = con_mtime(escribir(tmp_path, []), T0)
    con_mtime(escribir_xlsx(tmp_path, [["S003", 4400123, "ABCU1000003", "B", None, None]]), T0 + 10)
    correcciones.resolver_excel_pendiente(ruta_csv, tmp_path / "b", SEP, ENC, "%d/%m/%Y", [])
    assert ruta_csv.read_text(encoding=ENC).splitlines()[1].split(SEP)[1] == "4400123"


def test_xlsx_corrupto_no_tumba_la_corrida(tmp_path):
    ruta_csv = con_mtime(escribir(tmp_path, [["S004", "1", "ABCU1000004", "C", "", ""]]), T0)
    ruta_xlsx = tmp_path / "01-TEST.xlsx"
    ruta_xlsx.write_bytes(b"esto no es un archivo de Excel")
    con_mtime(ruta_xlsx, T0 + 10)
    antes, avisos = ruta_csv.read_bytes(), []
    assert not correcciones.resolver_excel_pendiente(ruta_csv, tmp_path / "b", SEP, ENC, "%d/%m/%Y",
                                                     avisos)
    assert ruta_csv.read_bytes() == antes and tipos(avisos) == ["correccion_excel_error"]
